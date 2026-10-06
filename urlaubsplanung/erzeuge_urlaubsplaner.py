"""Erzeugt Urlaubsplaner_Kreativtherapie.xlsx.

Aufruf:  python3 erzeuge_urlaubsplaner.py [Startdatum JJJJ-MM-TT]
Ohne Argument beginnt der Planungszeitraum am 01.01.2027.

Aufbau und Entscheidungen sind in README.md beschrieben.
"""
import sys
from datetime import date, timedelta

from dateutil.easter import easter
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation

START = date.fromisoformat(sys.argv[1]) if len(sys.argv) > 1 else date(2027, 1, 1)
TAGE = 366  # Planungszeitraum in Tagen (deckt auch Schaltjahre ab)

# Kapazitäten (bei Bedarf hier erhöhen und Skript neu ausführen)
N_PERS = 60        # Zeilen für Therapeut:innen
N_FACH = 10        # Fachrichtungen
N_STAT = 30        # Stationen
N_FEIER = 34       # Feiertage / Schließtage
N_URL = 600        # Urlaubseinträge
N_ST_PRO_PERS = 4  # Stationen je Person

# ---------- Layout-Konstanten ----------
P0 = 5                       # erste Personenzeile (Team, Kalender, Konflikt)
P1 = P0 + N_PERS - 1
F0 = P1 + 4                  # erste Fachrichtungs-Zeile im Kalender
F1 = F0 + N_FACH - 1
S0 = F1 + 4                  # erste Stations-Zeile im Kalender
S1 = S0 + N_STAT - 1
D0 = 7                       # erste Tages-Spalte (G)
D1 = D0 + TAGE - 1
DC0, DC1 = col(D0), col(D1)
U0, U1 = 5, 5 + N_URL - 1    # Urlaubseinträge

E_START, E_ENDE = "Einstellungen!$B$3", "Einstellungen!$B$4"
E_FEIER = f"Einstellungen!$I$7:$I${6 + N_FEIER}"
TEAM_NAMEN = f"Team!$A${P0}:$A${P1}"

# ---------- Stil ----------
FONT = "Arial"
f_norm = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=14, bold=True)
f_small = Font(name=FONT, size=8)
f_input = Font(name=FONT, size=10, color="0000FF")
fill_head = PatternFill("solid", fgColor="2F5496")
fill_input = PatternFill("solid", fgColor="FFF2CC")
fill_sec = PatternFill("solid", fgColor="D9E1F2")
fill_ex = PatternFill("solid", fgColor="EDEDED")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")
wrap = Alignment(wrap_text=True, vertical="top")

RED = PatternFill("solid", fgColor="F8696B", bgColor="F8696B")
YEL = PatternFill("solid", fgColor="FFEB84", bgColor="FFEB84")
GRN = PatternFill("solid", fgColor="C6EFCE", bgColor="C6EFCE")
BLU = PatternFill("solid", fgColor="8EA9DB", bgColor="8EA9DB")
GRY = PatternFill("solid", fgColor="E7E6E6", bgColor="E7E6E6")


def head(ws, row, labels, start_col=1, widths=None):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=text)
        c.font, c.fill, c.alignment, c.border = f_head, fill_head, Alignment(
            horizontal="center", vertical="center", wrap_text=True), box
        if widths:
            ws.column_dimensions[col(start_col + i)].width = widths[i]


def input_cells(ws, rng):
    for row in ws[rng]:
        for c in row:
            c.fill, c.font, c.border = fill_input, f_input, box


def set_font(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != FONT:
                c.font = Font(name=FONT, size=c.font.size or 10, bold=c.font.bold,
                              color=c.font.color)


wb = Workbook()

# =====================================================================
# Anleitung
# =====================================================================
wa = wb.active
wa.title = "Anleitung"
wa.column_dimensions["A"].width = 115
texte = [
    ("Urlaubsplaner Kreativtherapie-Team", f_title),
    ("Gemeinsame Übersicht der Urlaubswünsche. Das ist KEIN Genehmigungsverfahren: Die Datei zeigt, "
     "wann es im eigenen Fachteam oder auf einer Station eng wird, damit man sich rechtzeitig abspricht.", f_norm),
    ("", f_norm),
    ("SO TRAGE ICH MEINEN URLAUB EIN", f_bold),
    ("1. Blatt »Urlaub« öffnen, nächste freie Zeile suchen.", f_norm),
    ("2. Spalte A: eigenen Namen aus der Liste wählen (nicht tippen – die Liste kommt aus dem Blatt »Team«).", f_norm),
    ("3. Von / Bis als Datum eintragen (z. B. 12.07.2027). Ein einzelner Tag: Von = Bis.", f_norm),
    ("4. Status wählen: Wunsch = noch offen · abgesprochen = mit Kolleg:innen geklärt · eingereicht = offiziell beantragt.", f_norm),
    ("5. Spalten F–J füllen sich selbst. Steht in »Hinweis« ›Engpass – bitte absprechen‹, sind im Zeitraum an "
     "mindestens einem Tag mehr Personen weg als erlaubt (Spalte H = im Fachteam, Spalte I = auf einer meiner Stationen).", f_norm),
    ("6. Bei Engpass: im Blatt »Kalender« nachsehen, wer betroffen ist (rote Felder), und sich absprechen. "
     "Danach Zeitraum anpassen oder in »Bemerkung« die Absprache notieren.", f_norm),
    ("7. Urlaub verschoben? Einfach Von/Bis in der eigenen Zeile ändern. Urlaub fällt weg? Zeile leeren (Inhalte löschen).", f_norm),
    ("8. Datei speichern und schließen (siehe »Ablage« unten).", f_norm),
    ("", f_norm),
    ("SO LESE ICH DEN KALENDER", f_bold),
    ("• Oben: eine Zeile pro Person, eine Spalte pro Tag. »U« in Blau = Urlaub. »U« in Rot = an diesem Tag ist "
     "das Fachteam oder eine Station dieser Person über der Grenze.", f_norm),
    ("• Darunter: »Im Urlaub je Fachrichtung« und »Im Urlaub je Station« – Zahl = wie viele an dem Tag weg sind. "
     "Grün = unter der Grenze, Gelb = Grenze erreicht (keiner mehr dazu), Rot = Grenze überschritten.", f_norm),
    ("• Spalte B in diesen Zeilen zeigt die Grenze (max. gleichzeitig im Urlaub), Spalte C die Teamgröße.", f_norm),
    ("• Graue Spalten = Wochenende oder Feiertag. Orange umrandet = heute.", f_norm),
    ("• Der Kalender ist schreibgeschützt (ohne Passwort), damit niemand versehentlich Formeln überschreibt. "
     "Einträge immer im Blatt »Urlaub« machen.", f_norm),
    ("", f_norm),
    ("FARBEN IN DEN EINGABEBLÄTTERN", f_bold),
    ("• Gelb hinterlegt, blaue Schrift = hier darf/soll eingetragen werden. Alles andere wird berechnet.", f_norm),
    ("• Grau hinterlegte Zeilen mit »Beispiel« sind Musterdaten: vor dem Start löschen (siehe unten).", f_norm),
    ("", f_norm),
    ("EINRICHTUNG (einmalig, durch die verantwortliche Person)", f_bold),
    ("1. Blatt »Einstellungen«: Startdatum des Planungszeitraums (B3), Fachrichtungen mit Grenze, Stationen mit Grenze, "
     "Feiertage. Bundesweite Feiertage sind eingetragen – landesspezifische (z. B. Fronleichnam, Reformationstag) und "
     "Schließtage bitte ergänzen.", f_norm),
    ("2. Blatt »Team«: alle Therapeut:innen mit Fachrichtung und bis zu 4 Stationen. Tipp: nach Fachrichtung sortiert "
     "eintragen, dann steht der Kalender gruppiert. Jeder Name nur einmal (sonst werden Urlaube doppelt gezählt).", f_norm),
    ("3. Beispieldaten löschen: in »Team« und »Urlaub« die grauen Beispielzeilen markieren → Inhalte löschen "
     "(Entf-Taste, NICHT ›Zeilen löschen‹ – sonst verschieben sich Bezüge). Danach die graue Füllung entfernen oder belassen.", f_norm),
    ("4. Grenze (max. gleichzeitig im Urlaub) leer lassen = keine Grenze für dieses Team / diese Station.", f_norm),
    ("5. Neues Jahr: Datei kopieren, in der Kopie Startdatum ändern, alte Urlaubseinträge löschen, Feiertage aktualisieren.", f_norm),
    ("", f_norm),
    ("ABLAGE UND GLEICHZEITIGES ARBEITEN", f_bold),
    ("• Beste Variante: Datei in Microsoft Teams / SharePoint / OneDrive der Klinik ablegen. Dann können mehrere Personen "
     "gleichzeitig eintragen (auch im Browser mit Excel für das Web).", f_norm),
    ("• Nur Netzlaufwerk vorhanden: Excel sperrt die Datei, solange jemand sie geöffnet hat – andere sehen dann "
     "›schreibgeschützt‹. Deshalb: öffnen, eintragen, speichern, sofort schließen.", f_norm),
    ("• Die Datei enthält personenbezogene Daten (Abwesenheiten). Nur im Klinik-System ablegen, Zugriff nur fürs Team. "
     "Ggf. mit Personalrat/Betriebsrat abstimmen, ob ein solcher gemeinsamer Plan abgesprochen werden muss.", f_norm),
    ("", f_norm),
    ("GRENZEN DER DATEI", f_bold),
    (f"Platz für {N_PERS} Personen, {N_FACH} Fachrichtungen, {N_STAT} Stationen, {N_URL} Urlaubseinträge, "
     f"{N_ST_PRO_PERS} Stationen je Person, {TAGE} Tage. Mehr nötig? Generator-Skript (erzeuge_urlaubsplaner.py) mit "
     "größeren Werten neu ausführen und Daten hinüberkopieren – siehe README.md.", f_norm),
    ("Blatt »Konflikt« ist ein ausgeblendetes Hilfsblatt (Rechnung, welche Tage über der Grenze liegen). Nicht bearbeiten.", f_norm),
]
for i, (t, f) in enumerate(texte, start=1):
    c = wa.cell(row=i, column=1, value=t)
    c.font, c.alignment = f, Alignment(wrap_text=True, vertical="top")

# =====================================================================
# Einstellungen
# =====================================================================
we = wb.create_sheet("Einstellungen")
we["A1"] = "Einstellungen"
we["A1"].font = f_title
we["A3"], we["A4"] = "Startdatum Planungszeitraum", "Letzter Tag (berechnet)"
we["B3"] = START
we["B4"] = f"=B3+{TAGE - 1}"
for c in ("B3", "B4"):
    we[c].number_format = "DD.MM.YYYY"
input_cells(we, "B3:B3")
we["C3"] = "← eintragen. Der Kalender zeigt ab hier 366 Tage."
we["C3"].font = f_small
for c in ("A3", "A4"):
    we[c].font = f_bold
we["B4"].border = box
we.column_dimensions["A"].width = 26
we.column_dimensions["B"].width = 14

head(we, 6, ["Fachrichtung", "Max. gleichzeitig\nim Urlaub", "Personen\nim Team"], 1, [26, 14, 10])
head(we, 6, ["Station", "Max. gleichzeitig\nim Urlaub", "Therapeut:innen\nauf Station"], 5, [22, 14, 14])
head(we, 6, ["Feiertag / Schließtag", "Bezeichnung"], 9, [16, 26])
we.row_dimensions[6].height = 30
we.column_dimensions["D"].width = 3
we.column_dimensions["H"].width = 3

fach = [("Ergotherapie", 3), ("Musiktherapie", 2), ("Bewegungstherapie", 2), ("Kunsttherapie", 1)]
for i in range(N_FACH):
    r = 7 + i
    if i < len(fach):
        we.cell(r, 1, fach[i][0])
        we.cell(r, 2, fach[i][1])
    we.cell(r, 3, f'=IF(A{r}="","",COUNTIF(Team!$B${P0}:$B${P1},A{r}))').border = box
input_cells(we, f"A7:B{6 + N_FACH}")

st = [("Station 1A", 2), ("Station 1B", 1), ("Station 2A", 2), ("Station 3", 1), ("Tagesklinik", 1)]
st_cols = f"Team!$C${P0}:${col(2 + N_ST_PRO_PERS)}${P1}"
for i in range(N_STAT):
    r = 7 + i
    if i < len(st):
        we.cell(r, 5, st[i][0])
        we.cell(r, 6, st[i][1])
    we.cell(r, 7, f'=IF(E{r}="","",COUNTIF({st_cols},E{r}))').border = box
input_cells(we, f"E7:F{6 + N_STAT}")

o = easter(START.year)
feiertage = [
    (date(START.year, 1, 1), "Neujahr"),
    (o - timedelta(days=2), "Karfreitag"),
    (o + timedelta(days=1), "Ostermontag"),
    (date(START.year, 5, 1), "Tag der Arbeit"),
    (o + timedelta(days=39), "Christi Himmelfahrt"),
    (o + timedelta(days=50), "Pfingstmontag"),
    (date(START.year, 10, 3), "Tag der Deutschen Einheit"),
    (date(START.year, 12, 25), "1. Weihnachtstag"),
    (date(START.year, 12, 26), "2. Weihnachtstag"),
]
for i, (d, n) in enumerate(feiertage):
    we.cell(7 + i, 9, d)
    we.cell(7 + i, 10, n)
for r in range(7, 7 + N_FEIER):
    we.cell(r, 9).number_format = "DD.MM.YYYY"
input_cells(we, f"I7:J{6 + N_FEIER}")
we["I5"] = "Bundesweite Feiertage vorbelegt – Landesfeiertage ergänzen!"
we["I5"].font = f_small
we["B6"].comment = Comment("Leer lassen = keine Grenze. Es zählt jede Person, die an dem Tag "
                           "einen Eintrag im Blatt »Urlaub« hat (egal welcher Status).", "Urlaubsplaner")
we.freeze_panes = "A7"

dv_num = DataValidation(type="whole", operator="between", formula1="0", formula2="99",
                        showErrorMessage=True, error="Bitte eine ganze Zahl (0–99) eintragen oder leer lassen.")
we.add_data_validation(dv_num)
dv_num.add(f"B7:B{6 + N_FACH}")
dv_num.add(f"F7:F{6 + N_STAT}")

# =====================================================================
# Team
# =====================================================================
wt = wb.create_sheet("Team")
wt["A1"] = "Team – Stammdaten"
wt["A1"].font = f_title
wt["A2"] = ("Eine Zeile pro Person. Name eindeutig (z. B. »Anna M.«). Fachrichtung und Stationen aus der Liste wählen; "
            "Listen pflegt man im Blatt »Einstellungen«. Am besten nach Fachrichtung sortiert eintragen.")
wt["A2"].font = f_small
head(wt, 4, ["Name", "Fachrichtung"] + [f"Station {i + 1}" for i in range(N_ST_PRO_PERS)] + ["Bemerkung (z. B. Teilzeit)"],
     1, [22, 20] + [15] * N_ST_PRO_PERS + [30])
input_cells(wt, f"A{P0}:{col(3 + N_ST_PRO_PERS)}{P1}")

beispiel_team = [
    ("Beispiel Anna", "Ergotherapie", ["Station 1A", "Station 1B"]),
    ("Beispiel Ben", "Ergotherapie", ["Station 2A"]),
    ("Beispiel Clara", "Ergotherapie", ["Station 3", "Tagesklinik"]),
    ("Beispiel David", "Musiktherapie", ["Station 1A", "Station 2A"]),
    ("Beispiel Eva", "Musiktherapie", ["Station 1B", "Station 3"]),
    ("Beispiel Finn", "Bewegungstherapie", ["Station 1A", "Tagesklinik"]),
    ("Beispiel Gül", "Bewegungstherapie", ["Station 1B", "Station 2A", "Station 3"]),
]
for i, (n, fr, sts) in enumerate(beispiel_team):
    r = P0 + i
    wt.cell(r, 1, n)
    wt.cell(r, 2, fr)
    for j, s in enumerate(sts):
        wt.cell(r, 3 + j, s)
    for cc in range(1, 4 + N_ST_PRO_PERS):
        wt.cell(r, cc).fill = fill_ex

dv_fach = DataValidation(type="list", formula1=f"=Einstellungen!$A$7:$A${6 + N_FACH}", allow_blank=True)
dv_stat = DataValidation(type="list", formula1=f"=Einstellungen!$E$7:$E${6 + N_STAT}", allow_blank=True)
wt.add_data_validation(dv_fach)
wt.add_data_validation(dv_stat)
dv_fach.add(f"B{P0}:B{P1}")
dv_stat.add(f"C{P0}:{col(2 + N_ST_PRO_PERS)}{P1}")
# doppelte Namen rot markieren
wt.conditional_formatting.add(f"A{P0}:A{P1}", FormulaRule(
    formula=[f'AND(A{P0}<>"",COUNTIF($A${P0}:$A${P1},A{P0})>1)'], fill=RED))
wt.freeze_panes = f"A{P0}"

# =====================================================================
# Urlaub (Eingabe)
# =====================================================================
wu = wb.create_sheet("Urlaub")
wu["A1"] = "Urlaubswünsche – hier eintragen"
wu["A1"].font = f_title
wu["A2"] = ("Gelbe Spalten ausfüllen (A–E). F–J rechnen automatisch. »Engpass« heißt: an mindestens einem Tag im Zeitraum "
            "sind mehr Personen weg als erlaubt – bitte im Kalender nachsehen und absprechen.")
wu["A2"].font = f_small
head(wu, 4, ["Name", "Von", "Bis", "Status", "Bemerkung / Absprache", "Fachrichtung", "Arbeitstage\n(Mo–Fr ohne Feiertage)",
             "Tage über Grenze\nim Fachteam", "Tage über Grenze\nauf meinen Stationen", "Hinweis"],
     1, [22, 12, 12, 13, 34, 18, 13, 14, 16, 30])
wu.row_dimensions[4].height = 42
input_cells(wu, f"A{U0}:E{U1}")

beispiel_url = [
    ("Beispiel Anna", date(START.year, 7, 5), date(START.year, 7, 23), "abgesprochen", "mit Ben getauscht"),
    ("Beispiel Ben", date(START.year, 7, 26), date(START.year, 8, 13), "Wunsch", ""),
    ("Beispiel David", date(START.year, 7, 12), date(START.year, 7, 30), "Wunsch", ""),
    ("Beispiel Finn", date(START.year, 7, 19), date(START.year, 7, 30), "Wunsch", "Station 1A wäre dann leer – klären"),
    ("Beispiel Eva", date(START.year, 10, 11), date(START.year, 10, 15), "eingereicht", ""),
]
for i, rowv in enumerate(beispiel_url):
    for j, v in enumerate(rowv):
        wu.cell(U0 + i, 1 + j, v if v != "" else None)
    for cc in range(1, 6):
        wu.cell(U0 + i, cc).fill = fill_ex

for r in range(U0, U1 + 1):
    wu.cell(r, 2).number_format = "DD.MM.YYYY"
    wu.cell(r, 3).number_format = "DD.MM.YYYY"
    ok = f'OR($A{r}="",$B{r}="",$C{r}="",$C{r}<$B{r})'
    lo = f"MAX($B{r},{E_START})"
    hi = f"MIN($C{r},{E_ENDE})"
    rng = (f"OFFSET(Konflikt!${DC0}${P0 - 1},MATCH($A{r},{TEAM_NAMEN},0),{lo}-{E_START},1,{hi}-{lo}+1)")
    wu.cell(r, 6, f'=IF($A{r}="","",IFERROR(INDEX(Team!$B${P0}:$B${P1},MATCH($A{r},{TEAM_NAMEN},0)),"?"))')
    wu.cell(r, 7, f'=IF({ok},"",NETWORKDAYS($B{r},$C{r},{E_FEIER}))')
    wu.cell(r, 8, f'=IF({ok},"",IF({hi}<{lo},0,IFERROR(SUMPRODUCT(--(MOD({rng},2)=1)),"?")))')
    wu.cell(r, 9, f'=IF({ok},"",IF({hi}<{lo},0,IFERROR(SUMPRODUCT(--({rng}>=2)),"?")))')
    wu.cell(r, 10, (f'=IF($A{r}="","",IF(OR($B{r}="",$C{r}=""),"Von und Bis eintragen",'
                    f'IF($C{r}<$B{r},"Bis liegt vor Von",'
                    f'IF(ISNA(MATCH($A{r},{TEAM_NAMEN},0)),"Name fehlt im Blatt Team",'
                    f'IF(N($H{r})+N($I{r})>0,"Engpass – bitte absprechen",'
                    f'IF(OR($B{r}<{E_START},$C{r}>{E_ENDE}),"Teils außerhalb des Planungszeitraums","ok"))))))'))
    for cc in range(6, 11):
        wu.cell(r, cc).border = box
        wu.cell(r, cc).font = f_norm
    for cc in (7, 8, 9):
        wu.cell(r, cc).alignment = center

dv_name = DataValidation(type="list", formula1=f"={TEAM_NAMEN}", allow_blank=True, showErrorMessage=True,
                         error="Bitte den Namen aus der Liste wählen. Neue Personen zuerst im Blatt »Team« anlegen.")
dv_status = DataValidation(type="list", formula1='"Wunsch,abgesprochen,eingereicht"', allow_blank=True)
dv_date = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2100,12,31)",
                         showErrorMessage=True, error="Bitte ein Datum eintragen, z. B. 12.07.2027.")
dv_bis = DataValidation(type="custom", formula1=f"AND(ISNUMBER(C{U0}),C{U0}>=B{U0})", showErrorMessage=True,
                        error="»Bis« muss ein Datum sein und darf nicht vor »Von« liegen.")
for dv in (dv_name, dv_status, dv_date, dv_bis):
    wu.add_data_validation(dv)
dv_name.add(f"A{U0}:A{U1}")
dv_status.add(f"D{U0}:D{U1}")
dv_date.add(f"B{U0}:B{U1}")
dv_bis.add(f"C{U0}:C{U1}")

wu.conditional_formatting.add(f"J{U0}:J{U1}", FormulaRule(formula=[f'LEFT($J{U0},8)="Engpass "'], fill=RED))
wu.conditional_formatting.add(f"J{U0}:J{U1}", FormulaRule(formula=[f'$J{U0}="ok"'], fill=GRN))
wu.conditional_formatting.add(f"J{U0}:J{U1}", FormulaRule(
    formula=[f'AND($J{U0}<>"",$J{U0}<>"ok")'], fill=YEL))
wu.conditional_formatting.add(f"H{U0}:I{U1}", FormulaRule(formula=[f'N(H{U0})>0'], fill=RED))
wu.freeze_panes = f"B{U0}"
wu.auto_filter.ref = f"A4:J{U1}"

# =====================================================================
# Kalender
# =====================================================================
wk = wb.create_sheet("Kalender")
wk["A1"] = "Kalender – wer ist wann weg?"
wk["A1"].font = f_title
wk["A2"] = "Nur ansehen. Einträge im Blatt »Urlaub«."
wk["A2"].font = f_small
head(wk, 4, ["Name", "Fachrichtung"] + [f"St. {i + 1}" for i in range(N_ST_PRO_PERS)], 1, [22, 18] + [11] * N_ST_PRO_PERS)
wk["A3"] = "Monat / Tag →"
wk["A3"].font = f_small

MONATE = '"Jan","Feb","Mär","Apr","Mai","Jun","Jul","Aug","Sep","Okt","Nov","Dez"'
for d in range(D0, D1 + 1):
    L = col(d)
    wk.column_dimensions[L].width = 3.6
    wk.cell(2, d, f'=IF(OR(DAY({L}3)=1,COLUMN()={D0}),CHOOSE(MONTH({L}3),{MONATE})&" "&YEAR({L}3),"")').font = f_bold
    c3 = wk.cell(3, d, f"={E_START}+{d - D0}")
    c3.number_format = "DD"
    c3.font, c3.alignment = f_small, center
    c4 = wk.cell(4, d, f'=CHOOSE(WEEKDAY({L}3,2),"Mo","Di","Mi","Do","Fr","Sa","So")')
    c4.font, c4.alignment, c4.fill = Font(name=FONT, size=7, color="FFFFFF"), center, fill_head

# Personenzeilen
for r in range(P0, P1 + 1):
    wk.cell(r, 1, f'=IF(Team!A{r}="","",Team!A{r})')
    for cc in range(2, 3 + N_ST_PRO_PERS):
        wk.cell(r, cc, f'=IF(Team!{col(cc)}{r}="","",Team!{col(cc)}{r})').font = f_small
    for d in range(D0, D1 + 1):
        L = col(d)
        c = wk.cell(r, d, (f'=IF($A{r}="",0,IF(COUNTIFS(Urlaub!$A${U0}:$A${U1},$A{r},'
                           f'Urlaub!$B${U0}:$B${U1},"<="&{L}$3,Urlaub!$C${U0}:$C${U1},">="&{L}$3)>0,1,0))'))
        c.number_format = '"U";;;'
        c.alignment = center
        c.border = box

# Summen Fachrichtung
wk.cell(F0 - 1, 1, "Im Urlaub je Fachrichtung").font = f_bold
head(wk, F0 - 1, ["Grenze", "Pers."], 2)
for i in range(N_FACH):
    r, er = F0 + i, 7 + i
    wk.cell(r, 1, f'=IF(Einstellungen!A{er}="","",Einstellungen!A{er})').font = f_bold
    wk.cell(r, 2, f'=IF(Einstellungen!B{er}="",999,Einstellungen!B{er})').number_format = '[>=999]"–";0'
    wk.cell(r, 3, f'=Einstellungen!C{er}')
    for d in range(D0, D1 + 1):
        L = col(d)
        c = wk.cell(r, d, f'=IF($A{r}="",0,SUMIF($B${P0}:$B${P1},$A{r},{L}${P0}:{L}${P1}))')
        c.number_format, c.alignment, c.border = "0;;;", center, box

# Summen Station
wk.cell(S0 - 1, 1, "Im Urlaub je Station").font = f_bold
head(wk, S0 - 1, ["Grenze", "Pers."], 2)
for i in range(N_STAT):
    r, er = S0 + i, 7 + i
    wk.cell(r, 1, f'=IF(Einstellungen!E{er}="","",Einstellungen!E{er})').font = f_bold
    wk.cell(r, 2, f'=IF(Einstellungen!F{er}="",999,Einstellungen!F{er})').number_format = '[>=999]"–";0'
    wk.cell(r, 3, f'=Einstellungen!G{er}')
    match = "+".join(f"(${col(cc)}${P0}:${col(cc)}${P1}=$A{r})" for cc in range(3, 3 + N_ST_PRO_PERS))
    for d in range(D0, D1 + 1):
        L = col(d)
        c = wk.cell(r, d, f'=IF($A{r}="",0,SUMPRODUCT((({match})>0)*{L}${P0}:{L}${P1}))')
        c.number_format, c.alignment, c.border = "0;;;", center, box

grid = f"{DC0}{P0}:{DC1}{P1}"
sums = f"{DC0}{F0}:{DC1}{S1}"
cf = wk.conditional_formatting
cf.add(grid, FormulaRule(formula=[f"Konflikt!{DC0}{P0}>0"], fill=RED, stopIfTrue=True))
cf.add(grid, FormulaRule(formula=[f"{DC0}{P0}=1"], fill=BLU, font=Font(bold=True, color="FFFFFF"), stopIfTrue=True))
cf.add(sums, FormulaRule(formula=[f"{DC0}{F0}>$B{F0}"], fill=RED, stopIfTrue=True))
cf.add(sums, FormulaRule(formula=[f"AND({DC0}{F0}>0,{DC0}{F0}=$B{F0})"], fill=YEL, stopIfTrue=True))
cf.add(sums, FormulaRule(formula=[f"{DC0}{F0}>0"], fill=GRN, stopIfTrue=True))
wochenende = f"OR(WEEKDAY({DC0}$3,2)>5,COUNTIF({E_FEIER},{DC0}$3)>0)"
cf.add(f"{DC0}3:{DC1}{S1}", FormulaRule(formula=[wochenende], fill=GRY))
orange = Side(style="medium", color="ED7D31")
cf.add(f"{DC0}3:{DC1}{S1}", FormulaRule(formula=[f"{DC0}$3=TODAY()"], border=Border(left=orange, right=orange)))
monatsbeginn = Side(style="medium", color="2F5496")
cf.add(f"{DC0}2:{DC1}{S1}", FormulaRule(formula=[f"DAY({DC0}$3)=1"], border=Border(left=monatsbeginn)))

wk.freeze_panes = f"{DC0}{P0}"
wk.protection.sheet = True
wk.protection.formatColumns = False
wk.protection.formatRows = False
wk.protection.autoFilter = False
wk.protection.sort = False

# =====================================================================
# Konflikt (Hilfsblatt, ausgeblendet)
# 1 = Fachteam über Grenze, 2 = eine Station über Grenze, 3 = beides, 0 = kein Konflikt / nicht im Urlaub
# =====================================================================
wc = wb.create_sheet("Konflikt")
wc["A1"] = "Hilfsblatt: 1 = Fachteam über Grenze, 2 = Station über Grenze, 3 = beides. Nicht bearbeiten."
head(wc, 4, ["Name", "Idx Fach"] + [f"Idx St.{i + 1}" for i in range(N_ST_PRO_PERS)], 1)
for r in range(P0, P1 + 1):
    wc.cell(r, 1, f"=Kalender!A{r}")
    wc.cell(r, 2, f"=IFERROR(MATCH(Kalender!$B{r},Kalender!$A${F0}:$A${F1},0),0)")
    for cc in range(3, 3 + N_ST_PRO_PERS):
        wc.cell(r, cc, f'=IF(Kalender!{col(cc)}{r}="",0,IFERROR(MATCH(Kalender!{col(cc)}{r},Kalender!$A${S0}:$A${S1},0),0))')
    for d in range(D0, D1 + 1):
        L = col(d)
        fach_ = (f"AND($B{r}>0,INDEX(Kalender!{L}${F0}:{L}${F1},MAX(1,$B{r}))"
                 f">INDEX(Kalender!$B${F0}:$B${F1},MAX(1,$B{r})))")
        stat_ = ",".join(f"AND(${col(cc)}{r}>0,INDEX(Kalender!{L}${S0}:{L}${S1},MAX(1,${col(cc)}{r}))"
                         f">INDEX(Kalender!$B${S0}:$B${S1},MAX(1,${col(cc)}{r})))" for cc in range(3, 3 + N_ST_PRO_PERS))
        wc.cell(r, d, f"=IF(Kalender!{L}{r}<>1,0,IF({fach_},1,0)+IF(OR({stat_}),2,0))")
wc.sheet_state = "hidden"

for ws in wb.worksheets:
    set_font(ws)
wu.sheet_view.tabSelected = False
wb.active = wb.worksheets.index(wu)
wa.sheet_view.tabSelected = False
wu.sheet_view.tabSelected = True

out = "Urlaubsplaner_Kreativtherapie.xlsx"
wb.save(out)
print("gespeichert:", out)
